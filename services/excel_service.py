import io
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from database.queries import get_all_leads

async def generate_leads_excel() -> io.BytesIO:
    """
    Профессиональная двухстраничная выгрузка:
    Лист 1: Компактная сводная таблица (без простыней текста) с цветовыми статусами.
    Лист 2: Построчная детализация всех ответов брифа для удобной фильтрации.
    """
    leads = await get_all_leads(limit=500)
    
    wb = openpyxl.Workbook()
    
    # ---------------- ЛИСТ 1: СВОДКА ЗАЯВОК ----------------
    ws1 = wb.active
    ws1.title = "Сводка заявок"
    ws1.views.sheetView[0].showGridLines = True
    
    headers1 = [
        "Код", "Дата", "Клиент", "Телефон", 
        "Связь", "Удобное время", "Направление", 
        "Бюджет и сроки", "Статус CRM", "Ответственный"
    ]
    ws1.append(headers1)
    
    # Стили оформления
    header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    thin_border = Border(
        left=Side(style='thin', color='E0E0E0'),
        right=Side(style='thin', color='E0E0E0'),
        top=Side(style='thin', color='E0E0E0'),
        bottom=Side(style='thin', color='E0E0E0')
    )
    
    status_styles = {
        "new": (PatternFill(start_color="FFF2CC", fill_type="solid"), Font(color="7F6000", bold=True), "Новый"),
        "in_progress": (PatternFill(start_color="DDEBF7", fill_type="solid"), Font(color="1F4E79", bold=True), "В работе"),
        "meeting_scheduled": (PatternFill(start_color="E2EFDA", fill_type="solid"), Font(color="375623", bold=True), "Созвон назначен"),
        "deal_closed": (PatternFill(start_color="C6EFCE", fill_type="solid"), Font(color="006100", bold=True), "Договор подписан"),
        "rejected": (PatternFill(start_color="F2DCDB", fill_type="solid"), Font(color="9C0006", bold=True), "Отказ")
    }

    # Стилизуем шапку Листа 1
    for col_idx in range(1, len(headers1) + 1):
        cell = ws1.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = thin_border
    
    # ---------------- ЛИСТ 2: ДЕТАЛИЗАЦИЯ БРИФОВ ----------------
    ws2 = wb.create_sheet(title="Детализация брифов")
    ws2.views.sheetView[0].showGridLines = True
    headers2 = ["Код заявки", "Клиент", "Параметр брифа", "Ответ клиента"]
    ws2.append(headers2)
    
    header2_fill = PatternFill(start_color="2F5597", end_color="2F5597", fill_type="solid")
    for col_idx in range(1, len(headers2) + 1):
        cell = ws2.cell(row=1, column=col_idx)
        cell.fill = header2_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = thin_border

    row2_idx = 2

    # Заполнение данными
    for row1_idx, lead in enumerate(leads, start=2):
        code = lead.get("lead_code", "")
        client = lead.get("client_name", "")
        created_str = str(lead.get("created_at", ""))[:16] # Обрезаем секунды
        st_raw = lead.get("status", "new")
        st_fill, st_font, st_label = status_styles.get(st_raw, (None, None, st_raw))
        
        # Строка для Листа 1 (чистая, без простыней)
        ws1.append([
            code,
            created_str,
            client,
            lead.get("client_phone", ""),
            lead.get("messenger", ""),
            lead.get("preferred_time", ""),
            lead.get("branch", ""),
            lead.get("budget_estimate", ""),
            st_label,
            lead.get("manager_name", "") or "Не назначен"
        ])
        
        # Оформление ячеек Листа 1
        for col_idx in range(1, len(headers1) + 1):
            c = ws1.cell(row=row1_idx, column=col_idx)
            c.border = thin_border
            c.font = Font(name="Calibri", size=10)
            c.alignment = Alignment(vertical="center", horizontal="center" if col_idx in (1, 2, 4, 9) else "left")
            
        # Подсветка статуса
        status_cell = ws1.cell(row=row1_idx, column=9)
        if st_fill:
            status_cell.fill = st_fill
            status_cell.font = st_font

        # Заполнение Листа 2 (Построчный разбор ответов)
        raw_summary = lead.get("answers_summary", "")
        # Разбираем пункты вида "• Параметр: Значение"
        lines = [line.strip() for line in raw_summary.split("\n") if line.strip()]
        for line in lines:
            clean_line = line.replace("<b>", "").replace("</b>", "").replace("• ", "")
            if ":" in clean_line:
                param, val = clean_line.split(":", 1)
            else:
                param, val = "Пункт", clean_line
                
            ws2.append([code, client, param.strip(), val.strip()])
            for c_idx in range(1, 5):
                cell = ws2.cell(row=row2_idx, column=c_idx)
                cell.border = thin_border
                cell.font = Font(name="Calibri", size=10)
                cell.alignment = Alignment(vertical="center", horizontal="center" if c_idx == 1 else "left")
            row2_idx += 1

    # Настройка ширины колонок и фиксация шапок
    for ws, max_cols in [(ws1, len(headers1)), (ws2, len(headers2))]:
        ws.freeze_panes = "A2" # Шапка зафиксирована
        ws.row_dimensions[1].height = 25
        for col in ws.iter_cols(min_col=1, max_col=max_cols):
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = openpyxl.utils.get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

    # Сохраняем в буфер
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    return output